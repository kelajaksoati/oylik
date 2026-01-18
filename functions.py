from docx import Document
from fpdf import FPDF
import database
import os

def process_docx(file_path):
    doc = Document(file_path)
    for para in doc.paragraphs:
        if "|" in para.text:
            parts = para.text.split("|")
            database.add_question("Attestatsiya", parts[0].strip(), parts[1].strip(), parts[2].strip())

def create_pdf(user_name, results):
    if not os.path.exists("uploads/reports"): os.makedirs("uploads/reports")
    pdf = FPDF()
    pdf.add_page()
    pdf.add_font('Arial', '', 'assets/arial.ttf', uni=True)
    pdf.set_font("Arial", size=14)
    pdf.cell(200, 10, txt=f"Test Hisoboti: {user_name}", ln=True, align='C')
    pdf.set_font("Arial", size=10)
    for i, r in enumerate(results):
        status = "✅" if r['is_correct'] else "❌"
        pdf.multi_cell(0, 10, txt=f"{i+1}. {r['q']}\nSiz: {r['u']} | To'g'ri: {r['c']} {status}")
    path = f"uploads/reports/report_{user_name}.pdf"
    pdf.output(path)
    return path
